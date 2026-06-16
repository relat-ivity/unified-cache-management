# export VLLM_BENCH_BLOCK_SIZE=128
# python3 -m pytest /root/unified-cache-management/test/suites/E2E/test_vllm_bench.py -s
from __future__ import annotations

import csv
import importlib.util
import math
import os
import re
import shutil
import subprocess
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pytest


REPO_ROOT = Path(__file__).resolve().parents[3]
SCRIPT_PATH = REPO_ROOT / "examples" / "vllm_bench.sh"
DEFAULT_RESULT_DIR = REPO_ROOT / "result"
DEFAULT_PROFILE_LOG = REPO_ROOT / "deepseek_vllm.log"
DEFAULT_REPORT_DIR = REPO_ROOT / "test" / "results" / "vllm_bench_report"

RESULT_DIRS = {
    "CUDA": DEFAULT_RESULT_DIR,
    "GDR": DEFAULT_RESULT_DIR,
}
PROFILE_LOGS = {
    "CUDA": DEFAULT_PROFILE_LOG,
    "GDR": DEFAULT_PROFILE_LOG,
}
BACKENDS = ["CUDA", "GDR"]

SCRIPT_TIMEOUT_S = 21600

HIT_RATES = [100, 95, 90, 80, 50, 0]
LOAD_SINGLE_ROW = "\u5e73\u5747\u5355\u5361load\u65f6\u95f4 (ms)"
LOAD_CROSS_RANK_ROW = "\u5e73\u5747\u516b\u5361load\u65f6\u95f4 (ms)"
ROW_LABELS = [
    LOAD_SINGLE_ROW,
    LOAD_CROSS_RANK_ROW,
    "Mean TTFT (ms)",
    "Median TTFT (ms)",
    "P99 TTFT (ms)",
    "Mean TPOT (ms)",
    "Request throughput (req/s)",
    "Total token throughput (tok/s)",
]

BENCH_METRIC_PATTERNS = {
    "Mean TTFT (ms)": re.compile(r"Mean\s+TTFT\s+\(ms\):\s*([0-9.]+)", re.I),
    "Median TTFT (ms)": re.compile(r"Median\s+TTFT\s+\(ms\):\s*([0-9.]+)", re.I),
    "P99 TTFT (ms)": re.compile(r"P99\s+TTFT\s+\(ms\):\s*([0-9.]+)", re.I),
    "Mean TPOT (ms)": re.compile(r"Mean\s+TPOT\s+\(ms\):\s*([0-9.]+)", re.I),
    "Request throughput (req/s)": re.compile(
        r"Request\s+throughput\s+\(req/s\):\s*([0-9.]+)", re.I
    ),
    "Total token throughput (tok/s)": re.compile(
        r"Total\s+Token\s+throughput\s+\(tok/s\):\s*([0-9.]+)", re.I
    ),
}


@dataclass(frozen=True)
class BenchCase:
    hit_rate: int
    prefix_len: int
    log_name: str


BENCH_CASES = {
    100: BenchCase(100, 128000, "100.log"),
    95: BenchCase(95, 121600, "95.log"),
    90: BenchCase(90, 115200, "90.log"),
    80: BenchCase(80, 102400, "80.log"),
    50: BenchCase(50, 64000, "50.log"),
    0: BenchCase(0, 0, "0.log"),
}


def _block_sizes() -> list[int]:
    raw = os.environ.get("VLLM_BENCH_BLOCK_SIZE", "128")
    sizes = [int(item.strip()) for item in raw.split(",") if item.strip()]
    return sizes or [128]


def _run_bench_script() -> None:
    if not SCRIPT_PATH.exists():
        pytest.fail(f"vllm bench script not found: {SCRIPT_PATH}")

    bash_bin = shutil.which("bash")
    if not bash_bin:
        pytest.skip("bash not found")

    if not shutil.which("vllm"):
        pytest.skip("vllm CLI not found")

    DEFAULT_RESULT_DIR.mkdir(parents=True, exist_ok=True)

    proc = subprocess.run(
        [bash_bin, str(SCRIPT_PATH)],
        cwd=REPO_ROOT,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        timeout=SCRIPT_TIMEOUT_S,
    )
    run_log = DEFAULT_RESULT_DIR / "vllm_bench_script.log"
    run_log.write_text(proc.stdout, encoding="utf-8", errors="replace")
    assert proc.returncode == 0, f"vllm_bench.sh failed, see log: {run_log}"


def _parse_bench_log(path: Path) -> dict[str, float]:
    if not path.exists():
        return {}

    text = path.read_text(encoding="utf-8", errors="replace")
    metrics: dict[str, float] = {}
    for label, pattern in BENCH_METRIC_PATTERNS.items():
        matches = pattern.findall(text)
        if matches:
            metrics[label] = float(matches[-1])
    return metrics


def _load_fawa_analyzer():
    module_path = REPO_ROOT / "benchmarks" / "analyze_fawa_profile.py"
    spec = importlib.util.spec_from_file_location("analyze_fawa_profile", module_path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Failed to load {module_path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _transfer_metrics(
    profile_log: Path | None,
    hit_rate: int,
    prefix_len: int,
    block_size: int,
) -> dict[str, float]:
    if hit_rate == 0:
        return {
            LOAD_SINGLE_ROW: 0.0,
            LOAD_CROSS_RANK_ROW: 0.0,
        }
    if profile_log is None or not profile_log.exists():
        return {}

    hit_external = prefix_len // block_size
    analyzer = _load_fawa_analyzer()
    summary = analyzer.summarize_group("logs", [profile_log], hit_external)
    wall = summary["wall_ms"]
    cross_rank = summary["cross_rank_wall_ms"]
    return {
        LOAD_SINGLE_ROW: wall["mean"],
        LOAD_CROSS_RANK_ROW: cross_rank["mean"],
    }


def _fmt(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    if isinstance(value, float):
        return f"{value:.3f}".rstrip("0").rstrip(".")
    return str(value)


def _build_table(block_size: int) -> dict[str, dict[tuple[int, str], Any]]:
    table = {label: {} for label in ROW_LABELS}

    for hit_rate in HIT_RATES:
        case = BENCH_CASES[hit_rate]
        for backend in BACKENDS:
            result_dir = RESULT_DIRS[backend]
            profile_log = PROFILE_LOGS[backend]
            bench_metrics = _parse_bench_log(result_dir / case.log_name)
            transfer = _transfer_metrics(
                profile_log, hit_rate, case.prefix_len, block_size
            )

            for row in ROW_LABELS:
                value = transfer.get(row, bench_metrics.get(row))
                table[row][(hit_rate, backend)] = value

    return table


def _write_csv(
    table: dict[str, dict[tuple[int, str], Any]],
    path: Path,
) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(
            [""]
            + [f"\u590d\u7528{hit_rate}%" for hit_rate in HIT_RATES for _ in BACKENDS]
        )
        writer.writerow([""] + [backend for _ in HIT_RATES for backend in BACKENDS])
        for row in ROW_LABELS:
            writer.writerow(
                [row]
                + [
                    _fmt(table[row].get((hit_rate, backend)))
                    for hit_rate in HIT_RATES
                    for backend in BACKENDS
                ]
            )


def _write_markdown(
    table: dict[str, dict[tuple[int, str], Any]],
    path: Path,
) -> None:
    columns = [
        f"\u590d\u7528{hit_rate}% {backend}"
        for hit_rate in HIT_RATES
        for backend in BACKENDS
    ]
    lines = ["| \u6307\u6807 | " + " | ".join(columns) + " |"]
    lines.append("|" + "|".join(["---"] * (len(columns) + 1)) + "|")
    for row in ROW_LABELS:
        values = [
            _fmt(table[row].get((hit_rate, backend)))
            for hit_rate in HIT_RATES
            for backend in BACKENDS
        ]
        lines.append(f"| {row} | " + " | ".join(values) + " |")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _write_xlsx(
    table: dict[str, dict[tuple[int, str], Any]],
    path: Path,
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "vllm_bench"
    ws.cell(row=1, column=1, value="")
    ws.cell(row=2, column=1, value="")

    col = 2
    for hit_rate in HIT_RATES:
        start = col
        for backend in BACKENDS:
            ws.cell(row=2, column=col, value=backend)
            if backend == "GDR":
                ws.cell(row=2, column=col).fill = PatternFill(
                    "solid", fgColor="E2F0D9"
                )
            col += 1
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=col - 1)
        ws.cell(row=1, column=start, value=f"\u590d\u7528{hit_rate}%")

    for row_idx, row_label in enumerate(ROW_LABELS, start=3):
        ws.cell(row=row_idx, column=1, value=row_label)
        ws.cell(row=row_idx, column=1).font = Font(bold=True)
        col = 2
        for hit_rate in HIT_RATES:
            for backend in BACKENDS:
                value = table[row_label].get((hit_rate, backend))
                ws.cell(row=row_idx, column=col, value=None if _fmt(value) == "" else value)
                if backend == "GDR":
                    ws.cell(row=row_idx, column=col).fill = PatternFill(
                        "solid", fgColor="E2F0D9"
                    )
                col += 1

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")
    for cell in ws["A"]:
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.column_dimensions["A"].width = 28
    for col_idx in range(2, 2 + len(HIT_RATES) * len(BACKENDS)):
        ws.column_dimensions[ws.cell(row=2, column=col_idx).column_letter].width = 14
    wb.save(path)


@pytest.mark.stage(2)
@pytest.mark.feature("vllm_bench_report")
def test_vllm_bench_script_and_report():
    _run_bench_script()

    DEFAULT_REPORT_DIR.mkdir(parents=True, exist_ok=True)

    for block_size in _block_sizes():
        table = _build_table(block_size)
        found_bench_metric = any(
            table[row].get((hit_rate, backend)) is not None
            for row in BENCH_METRIC_PATTERNS
            for hit_rate in HIT_RATES
            for backend in BACKENDS
        )
        assert found_bench_metric, (
            "No vllm bench metrics parsed. Check result logs under: "
            + ", ".join(str(path) for path in set(RESULT_DIRS.values()))
        )

        _write_csv(
            table, DEFAULT_REPORT_DIR / f"vllm_bench_summary_bs{block_size}.csv"
        )
        _write_markdown(
            table, DEFAULT_REPORT_DIR / f"vllm_bench_summary_bs{block_size}.md"
        )
        _write_xlsx(
            table, DEFAULT_REPORT_DIR / f"vllm_bench_summary_bs{block_size}.xlsx"
        )

    print(f"[INFO] vLLM bench summary written to: {DEFAULT_REPORT_DIR}")
